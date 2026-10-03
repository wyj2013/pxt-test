from pathlib import Path
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\物聯網通訊實務")
ASSETS = ROOT / "assets"
ASSETS.mkdir(exist_ok=True)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def note(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "FFF2CC")
    p = cell.paragraphs[0]
    r = p.add_run(f"重點註解｜{title}\n")
    r.bold = True
    r.font.color.rgb = RGBColor(156, 101, 0)
    p.add_run(text)
    doc.add_paragraph()

def picture(doc, path, width=6.2, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].italic = True
        c.runs[0].font.size = Pt(9)

def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(t.rows[0].cells, headers):
        cell_text(c, h, True, (255, 255, 255)); shade(c, '168AAD')
    for row in rows:
        cells = t.add_row().cells
        for c, value in zip(cells, row): cell_text(c, value)
    return t

# Copy the user's photos to stable filenames.
for i, src in enumerate(sorted((ROOT / "image").glob("*.png")), 1):
    shutil.copy2(src, ASSETS / f"photo_{i}.png")

arch = ASSETS / "system_architecture.png"
flow = ASSETS / "data_flow.png"
shutil.copy2(Path(r"C:\Users\PC2-38\.codex\generated_images\01a0ffff-aa2f-7fc2-92e3-7d2e8addab2f\exec-c755ad01-170f-4699-96f5-6658de57104f.png"), arch)
shutil.copy2(Path(r"C:\Users\PC2-38\.codex\generated_images\01a0ffff-aa2f-7fc2-92e3-7d2e8addab2f\exec-c39ad3ee-d7be-482e-b03b-87aeb4833fb1.png"), flow)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Microsoft JhengHei'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
doc.styles['Normal'].font.size = Pt(10.5)

# Cover page
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('學習歷程\n'); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor(0, 102, 153)
r = p.add_run('20261003物聯網通訊實務研習'); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = RGBColor(0, 150, 150)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('\n主題：Arduino ESP8266 × PMS5003 × MQTT 空氣品質物聯網實作\n').font.size = Pt(14)
p.add_run('成果版本：PMS5003 感測資料每 10 秒傳送至 mqttgo.io').font.size = Pt(11)
doc.add_paragraph('\n')
table(doc, ['項目', '內容'], [
    ('研習日期', '2026 年 10 月 3 日'), ('主控板', 'HW628 / ESP8266'),
    ('感測器', 'Plantower PMS5003'), ('MQTT 主題', 'phmhs/aqi')])
doc.add_page_break()

doc.add_heading('一、學習目標', level=1)
for x in ['熟悉 Arduino IDE、Arduino CLI、ESP8266 核心與 CH340 驅動的安裝與檢查。', '理解 ESP8266 透過 UART 讀取 PMS5003 空氣懸浮微粒資料。', '建立 Wi‑Fi 與 MQTT 通訊，將感測資料發布到指定主題。', '練習由序列埠、編譯結果與 MQTT 訊息進行系統化除錯。']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('二、開發環境與成果', level=1)
table(doc, ['項目', '版本／設定', '檢查結果'], [
    ('Arduino IDE', '2.3.10', '已安裝'), ('Arduino CLI', '1.5.1', '可編譯與上傳'),
    ('ESP8266 核心', '3.0.2', '已安裝'), ('CH340 驅動', '3.5.2019.1', '正常'),
    ('開發板連接埠', 'COM4', '可燒錄 ESP8266')])
note(doc, 'COM4 的判讀', '系統辨識到 USB-SERIAL CH340（VID_1A86、PID_7523），因此 Arduino IDE 與 Arduino CLI 都使用 COM4。')

doc.add_heading('三、系統架構', level=1)
doc.add_paragraph('PMS5003 量測 PM1.0、PM2.5、PM10，ESP8266 透過 Wi‑Fi 連到 MQTT Broker，再由訂閱端取得資料。')
picture(doc, arch, 6.5, '圖 1　ESP8266 空氣品質 MQTT 系統架構圖（imagegen 製作）')
note(doc, '感測器能力', '本專案使用 PMS5003，不是 PMS5003T，因此沒有溫度與濕度感測功能；temperature 與 humidity 以 -1 表示不適用。')

doc.add_heading('四、實體接線與成果照片', level=1)
doc.add_paragraph('PMS5003 使用 5V 供電，UART 資料線採交叉連接。')
table(doc, ['PMS5003', 'HW628 / ESP8266'], [
    ('VCC', '5V / VIN / VU'), ('GND', 'GND'),
    ('TXD', 'D6 / GPIO12（ESP8266 RX）'), ('RXD', 'D5 / GPIO14（ESP8266 TX，可選）')])
note(doc, 'UART 方向', '感測器 TXD 接 ESP8266 的 RX（D6）；感測器 RXD 接 ESP8266 的 TX（D5）。只讀取資料時，RXD 可不接。')
picture(doc, ASSETS / 'photo_1.png', 4.8, '圖 2　HW628、PMS5003 與 CH340 接線實體照片')
picture(doc, ASSETS / 'photo_3.png', 5.0, '圖 3　PMS5003 與 ESP8266 實作環境照片')

doc.add_heading('五、程式與資料格式', level=1)
doc.add_paragraph('專案檔案分為兩個版本：')
doc.add_paragraph('LED_01\\LED_01.ino：內建 LED 每 0.5 秒閃爍，作為硬體與燒錄測試。', style='List Bullet')
doc.add_paragraph('mqtt_01\\mqtt_01.ino：讀取 PMS5003，連線 Wi‑Fi 與 MQTT，每 10 秒發布一次。', style='List Bullet')
doc.add_paragraph('最終 MQTT 設定：')
table(doc, ['設定項目', '內容'], [('Wi‑Fi SSID', 'PHMHSCS02'), ('MQTT Server', 'mqttgo.io'), ('Port', '1883'), ('Topic', 'phmhs/aqi'), ('週期', '每 10 秒'), ('資料格式', 'JSON')])
doc.add_paragraph('有效資料範例：')
code = doc.add_paragraph(); r = code.add_run('{"pm01":10,"pm25":15,"pm10":15,"temperature":-1.0,"humidity":-1.0}')
r.font.name = 'Consolas'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0, 90, 0)
note(doc, '錯誤資料策略', '若 PMS5003 沒有收到有效封包，程式仍會發布 JSON，但欄位使用 -1，避免訂閱端誤認為資料遺失。')

doc.add_heading('六、資料流程', level=1)
picture(doc, flow, 5.3, '圖 4　PMS5003 到 MQTT 的資料流程圖（imagegen 製作）')

doc.add_heading('七、除錯過程紀錄', level=1)
steps = [
    ('1. 開發環境檢查', '完成 Arduino IDE 2.3.10、CLI 1.5.1、ESP8266 核心 3.0.2 與 CH340 驅動確認。'),
    ('2. COM Port 確認', '插入開發板後辨識到 USB-SERIAL CH340（COM4），使用 COM4 完成燒錄。'),
    ('3. Wi‑Fi 連線', '序列埠曾長時間顯示 Connecting to WiFi；重試後成功取得 192.168.1.116，後續 DHCP 也曾分配 192.168.1.117。'),
    ('4. MQTT 驗證', '成功連線 mqttgo.io:1883，主題 phmhs/aqi 可看到 Published to phmhs/aqi。'),
    ('5. 型號修正', '一開始誤以為是 PMS5003T，後來確認實際型號為 PMS5003，改用 PMS5003 的被動讀取命令與 32 位元組封包。'),
    ('6. 最終驗證', 'PMS5003 成功回傳 PM1.0、PM2.5、PM10，例如 10、15、15；溫度與濕度固定為 -1。')]
for title, body in steps:
    p = doc.add_paragraph(); r = p.add_run(title + '：'); r.bold = True; r.font.color.rgb = RGBColor(0, 102, 153); p.add_run(body)
picture(doc, ASSETS / 'photo_2.png', 5.2, '圖 5　MQTT／空氣品質資料畫面成果照片')

doc.add_heading('八、學習反思', level=1)
for x in [
    '本次實作讓我從能燒錄 LED，逐步完成感測器資料經 Wi‑Fi 傳送到 MQTT 的完整物聯網流程。',
    '我學到除錯必須分層確認：先確認 COM4 與燒錄，再看 Wi‑Fi IP，接著確認 MQTT，最後才判斷感測器封包。',
    'PMS5003 與 PMS5003T 的差異非常重要；PMS5003 沒有溫溼度功能，不能只修改欄位位置，必須依實際型號選擇正確封包格式。',
    '未來若要加入溫溼度，應另接 BME280 或 SHT30 等環境感測器，而不是從 PMS5003 推算。',
    '後續可加入 MQTT 帳號密碼與 TLS、Wi‑Fi 逾時重試、感測器錯誤碼與長期趨勢圖。']:
    doc.add_paragraph(x, style='List Bullet')

doc.add_heading('九、結論', level=1)
doc.add_paragraph('本次研習完成 Arduino ESP8266 開發環境建置、HW628 與 PMS5003 接線、資料讀取、Wi‑Fi 連線、MQTT 發布與錯誤值處理。最終系統能以 COM4 燒錄並運作，透過 phmhs/aqi 每 10 秒傳送空氣品質資料。')
doc.add_heading('附錄：重要檔案', level=1)
for x in ['LED_01\\LED_01.ino：LED 閃爍測試程式', 'mqtt_01\\mqtt_01.ino：PMS5003 + Wi‑Fi + MQTT 完整程式', 'assets\\system_architecture.png：系統架構圖', 'assets\\data_flow.png：資料流程圖']:
    doc.add_paragraph(x, style='List Bullet')

out = ROOT / '20261003物聯網通訊實務研習.docx'
doc.save(out)
print(out)
