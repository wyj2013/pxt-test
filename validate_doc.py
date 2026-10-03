from docx import Document
doc = Document(r"D:\物聯網通訊實務\20261003物聯網通訊實務研習.docx")
print(len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes))
